from pathlib import Path

from openpyxl import load_workbook

from app.reports.exporter import export_selection_report
from app.tools.selection_decision import Recommendation, SelectionDecision
from app.tools.shopping_summary import ShoppingSummaryOutput, _to_report


def _summary() -> ShoppingSummaryOutput:
    decision = SelectionDecision(
        product_id="A1",
        title="=unsafe formula title",
        platform="amazon",
        selling_price_cny=500,
        landed_cost_cny=180,
        total_cost_cny=250,
        net_profit_cny=250,
        profit_margin=0.5,
        roi=1.0,
        supplier_risk_score=0.1,
        overall_score=0.91,
        confidence=1.0,
        recommendation=Recommendation.RECOMMEND,
        reasons=["利润能力达到目标"],
        risks=[],
        missing_data=[],
    )
    return ShoppingSummaryOutput(
        final_text="推荐进入小批量验证。",
        decisions=[decision],
        report=[_to_report(decision)],
        learned_preferences=[],
    )


def test_export_selection_report_writes_pdf_and_xlsx(tmp_path: Path) -> None:
    exported = export_selection_report(_summary(), tmp_path, "selection-report")

    assert exported.pdf_path.is_absolute()
    assert exported.xlsx_path.is_absolute()
    assert exported.pdf_path.read_bytes().startswith(b"%PDF")
    pdf_bytes = exported.pdf_path.read_bytes()
    assert b"/Type /Page" in pdf_bytes
    assert b"NotoSansSC" in pdf_bytes
    assert b"/FontFile2" in pdf_bytes

    workbook = load_workbook(exported.xlsx_path, data_only=False)
    sheet = workbook["Selection Report"]
    assert sheet["A1"].value == "product_id"
    assert sheet["A2"].value == "A1"
    assert sheet["B2"].value == "'=unsafe formula title"
    assert sheet["I2"].value == 0.5


def test_export_rejects_unsafe_basename(tmp_path: Path) -> None:
    try:
        export_selection_report(_summary(), tmp_path, "../escape")
    except ValueError as exc:
        assert "basename" in str(exc)
    else:
        raise AssertionError("unsafe basename must be rejected")


def test_export_includes_evidence_sheet(tmp_path) -> None:
    """XLSX 证据页与 PDF 证据段应携带每个数字的口径。"""
    from openpyxl import load_workbook

    from app.evidence.models import EvidenceItem, EvidenceMethod
    from app.reports.exporter import export_selection_report
    from app.tools.selection_decision import (
        Recommendation,
        SelectionDecision,
    )
    from app.tools.shopping_summary import ShoppingSummaryOutput, _to_report

    decision = SelectionDecision(
        product_id="A1",
        title="测试商品",
        platform="amazon",
        overall_score=0.8,
        confidence=0.9,
        recommendation=Recommendation.RECOMMEND,
        evidence=[
            EvidenceItem(
                field="procurement_cost_cny",
                value=12.5,
                method=EvidenceMethod.EXTRACTED,
                source="web_search",
                detail="1688 搜索结果 | https://example.com/1",
                confidence=0.7,
                observed_at="2026-10-09T00:00:00Z",
                tool="procurement_quote",
            )
        ],
    )
    summary = ShoppingSummaryOutput(
        final_text="采购价 12.5 元（来自 1688 搜索证据）",
        decisions=[decision],
        report=[_to_report(decision)],
        learned_preferences=[],
    )
    result = export_selection_report(summary, tmp_path)
    workbook = load_workbook(result.xlsx_path)
    assert "证据" in workbook.sheetnames
    sheet = workbook["证据"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][1] == "field"
    assert rows[1][1] == "procurement_cost_cny"
    assert rows[1][3] == "提取"
    assert rows[1][5] == 0.7
